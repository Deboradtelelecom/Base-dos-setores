import openpyxl, pickle
from collections import defaultdict

agg = pickle.load(open('/tmp/agg_combustivel_v5_final.pkl', 'rb'))

MONTHS = ['MAIO/2026', 'JUNHO/2026', 'JULHO/2026', 'AGOSTO/2026']
MONTH_COL_IDX = {'MAIO/2026': 9, 'JUNHO/2026': 10, 'JULHO/2026': 11, 'AGOSTO/2026': 12}
CATS = ('Combustível', 'Manutenção de Veículos', 'Rastreamento de veículos', 'Licenciamento e IPVA')

FN = 'Custos_Regionais_DTEL.xlsx'
wb = openpyxl.load_workbook(FN, data_only=False)
ws = wb['Custo Regional']
det = wb['Detalhe']

rows_by_gestor_cat = {}
for r in range(2, ws.max_row + 1):
    desc = ws.cell(r, 5).value
    loc = ws.cell(r, 3).value
    gestor = ws.cell(r, 2).value
    if loc == 'Despesa Operacional' and desc in CATS:
        rows_by_gestor_cat[(gestor, desc)] = r

kept = []
for r in range(2, det.max_row + 1):
    vals = [det.cell(r, c).value for c in range(1, 8)]
    if vals[2] in CATS:
        continue
    if all(v is None for v in vals):
        continue
    kept.append(vals)

new_det_rows = []
for (gestor, cat), months in sorted(agg.items()):
    r = rows_by_gestor_cat.get((gestor, cat))
    chave = ws.cell(r, 15).value if r else f'{gestor.upper()}|DESPESA OPERACIONAL|{cat.upper()}||1'
    for m in MONTHS:
        v = round(months.get(m, 0.0), 2)
        if v:
            new_det_rows.append([gestor, 'Despesa Operacional', cat, m, 'Planilha Combustivel (Grupo Dtel: DTEL/GOONET/ORBIX/DMAIS/SPEED + recuperado via folha)', v, chave])

all_det_rows = kept + new_det_rows
max_r = max(det.max_row, len(all_det_rows)+1)
for r in range(2, max_r+1):
    for c in range(1, 8):
        det.cell(r, c).value = None
for i, row in enumerate(all_det_rows, start=2):
    for c, val in enumerate(row, start=1):
        det.cell(i, c, val)

atualizados = []
faltando = []
for (gestor, cat), months in sorted(agg.items()):
    r = rows_by_gestor_cat.get((gestor, cat))
    if r is None:
        faltando.append((gestor, cat, {m: round(v,2) for m,v in months.items()}))
        continue
    for m in MONTHS:
        v = round(months.get(m, 0.0), 2)
        ws.cell(r, MONTH_COL_IDX[m], v)
    ws.cell(r, 7, 'Planilha Combustivel: EMPRESA=Grupo Dtel (DTEL/DTEL TELECOM/GOONET/ORBIX/DMAIS/SPEED) + recuperado via cruzamento condutor->equipe da folha (25/09/2026, correção 2)')
    atualizados.append(r)

# zera categorias/meses de gestores que ficaram sem nenhum valor no novo agg (para nao deixar valor antigo)
gestores_cats_presentes = set(agg.keys())
for (gestor, cat), r in rows_by_gestor_cat.items():
    if (gestor, cat) not in gestores_cats_presentes:
        for m in MONTHS:
            ws.cell(r, MONTH_COL_IDX[m], 0)

print('Atualizados:', len(atualizados), 'Faltando linha:', faltando)
wb.save(FN)
print('Salvo.')
